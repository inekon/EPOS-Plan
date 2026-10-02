"""Tests für wsp/export_excel.py.

Grundsatz: Das Testprofil wird hier synthetisch konstruiert (Sinus-Jahresgang,
kein demandlib), die Auslegungsergebnisse werden aber **real** über
``wsp.analysis``, ``wsp.sizing_dhw`` und ``wsp.sizing_buffer`` berechnet.
Geprüft wird anschließend die wieder eingelesene Arbeitsmappe.
"""
from __future__ import annotations

import io
import os
import sys

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from wsp import analysis, export_excel  # noqa: E402
from wsp.models import (Building, BuildingCategory, BufferParams,  # noqa: E402
                        DHWParams, ProfileSet, WeatherConfig, project_to_dict)
from wsp.sizing_buffer import size_buffer  # noqa: E402
from wsp.sizing_dhw import size_dhw  # noqa: E402

JAHR = 2026          # kein Schaltjahr -> 8760 Stunden
N_WE = 12
PERS_JE_WE = 2.5
ERWARTETE_BLAETTER = ["Eingaben", "Lastgang", "JDL", "Tagesprofile",
                      "TWW-Auslegung", "Puffer-Auslegung", "Kennzahlen"]


# --------------------------------------------------------------------- Fixtures

def _profil(resolution: str = "h") -> ProfileSet:
    """Synthetisches Jahresprofil (kein demandlib).

    Q_heiz: Sinus-Jahresgang (Winterspitze) mit Tagesgang.
    Q_tww:  Grundlast plus Morgen-/Abendzapfspitze.
    """
    freq = resolution
    idx = pd.date_range(f"{JAHR}-01-01", f"{JAHR + 1}-01-01", freq=freq,
                        inclusive="left")
    tag = idx.dayofyear.to_numpy(dtype=float)
    stunde = (idx.hour + idx.minute / 60.0).to_numpy(dtype=float)

    # Jahresgang: Maximum am 1. Januar, Minimum Anfang Juli
    saison = 0.5 * (1.0 + np.cos(2.0 * np.pi * (tag - 1.0) / 365.0))
    tagesgang = 1.0 + 0.35 * np.cos(2.0 * np.pi * (stunde - 7.0) / 24.0)
    q_heiz = 8.0 + 92.0 * saison * tagesgang

    q_tww = np.full(len(idx), 3.0)
    q_tww[(stunde >= 6.0) & (stunde < 9.0)] = 45.0
    q_tww[(stunde >= 18.0) & (stunde < 21.0)] = 38.0

    df = pd.DataFrame({"Q_heiz": q_heiz, "Q_tww": q_tww}, index=idx)
    df.index.name = "Zeit"
    return ProfileSet(
        df=df,
        resolution=resolution,
        source="synthetisch",
        meta={
            "quelle": "Testprofil (Sinus-Jahresgang, konstruiert)",
            "jahr": JAHR,
            "n_we": N_WE,
            "try_region": 12,
            "try_region_name": "Region 12 (Test)",
        },
    )


def _projekt(dhw_params: DHWParams, buffer_params: BufferParams) -> dict:
    gebaeude = [
        Building(name="MFH Musterstraße 1", category=BuildingCategory.MFH,
                 n_we=N_WE, n_persons=PERS_JE_WE, q_heiz_a=180000.0,
                 q_tww_a=32000.0, copies=1),
        Building(name="Bürotrakt", category=BuildingCategory.GHD,
                 q_heiz_a=40000.0, bdew_branch="gko", tww_share=0.08),
    ]
    projekt = {
        "projektname": "Testprojekt Wärmespeicher",
        "datum": "2026-07-28",
        "bearbeiter": "pytest",
    }
    projekt.update(project_to_dict(gebaeude, WeatherConfig(try_region=12,
                                                           year=JAHR),
                                   dhw_params, buffer_params))
    return projekt


@pytest.fixture(scope="module")
def ergebnisse():
    """Reale Berechnung aller Ergebnisse für das Testprofil."""
    profile = _profil("h")
    kz = analysis.kennzahlen(profile)

    dhw_params = DHWParams(t_speicher=60.0, t_kalt=10.0, p_lade=25.0,
                           zirkulation_kw=1.5, zuschlag=0.15)
    buffer_params = BufferParams(p_wp=70.0, p_wp_min=20.0, dt_puffer=10.0,
                                 t_min_lauf_min=10.0, abtau_l_pro_kw=20.0,
                                 sperrzeiten=[(11, 2), (17, 2)],
                                 ziel_deckung=1.0)

    dhw = size_dhw(dhw_params, n_we=N_WE, persons_per_we=PERS_JE_WE,
                   profile_tww=profile.df["Q_tww"], dt_h=profile.dt_h)
    buf = size_buffer(buffer_params, profile_heiz=profile.df["Q_heiz"],
                      dt_h=profile.dt_h)

    return {
        "profile": profile,
        "kennzahlen": kz,
        "dhw": dhw,
        "buffer": buf,
        "projekt": _projekt(dhw_params, buffer_params),
    }


@pytest.fixture(scope="module")
def wb_bytesio(ergebnisse):
    """Bericht in einen BytesIO schreiben und wieder einlesen."""
    puffer = io.BytesIO()
    export_excel.write_report(puffer, ergebnisse["profile"],
                              ergebnisse["kennzahlen"], ergebnisse["dhw"],
                              ergebnisse["buffer"], ergebnisse["projekt"])
    assert puffer.tell() > 0, "In den BytesIO wurde nichts geschrieben"
    puffer.seek(0)
    return load_workbook(puffer)


# ----------------------------------------------------------------- Hilfsmittel

def _werte(ws) -> list:
    """Alle Zellwerte eines Blatts als flache Liste."""
    return [c for zeile in ws.iter_rows(values_only=True) for c in zeile
            if c is not None]


def _texte(ws) -> str:
    """Alle Textzellen eines Blatts zu einem String zusammengefasst."""
    return "\n".join(str(v) for v in _werte(ws))


def _enthaelt_zahl(ws, zahl: float, tol: float = 1e-6) -> bool:
    for v in _werte(ws):
        if isinstance(v, (int, float)) and not isinstance(v, bool):
            if abs(float(v) - float(zahl)) <= tol:
                return True
    return False


# --------------------------------------------------------------------- Tests

def test_alle_blaetter_vorhanden(wb_bytesio):
    assert wb_bytesio.sheetnames == ERWARTETE_BLAETTER


def test_lastgang_zeilenzahl_8760_plus_header(wb_bytesio):
    ws = wb_bytesio["Lastgang"]
    assert ws.max_row == 8761, f"erwartet 8760 Datenzeilen + Kopf, ist {ws.max_row}"
    assert ws.max_column == 4
    assert [c.value for c in ws[1]] == ["Datum/Zeit", "Q_heiz [kW]",
                                        "Q_tww [kW]", "Q_total [kW]"]
    assert ws["A1"].font.bold, "Kopfzeile muss fett sein"


def test_lastgang_werte_und_formate(wb_bytesio, ergebnisse):
    ws = wb_bytesio["Lastgang"]
    df = ergebnisse["profile"].df
    # erste Datenzeile
    assert ws.cell(row=2, column=1).value == df.index[0].to_pydatetime()
    assert ws.cell(row=2, column=2).value == pytest.approx(
        float(df["Q_heiz"].iloc[0]), rel=1e-9)
    assert ws.cell(row=2, column=4).value == pytest.approx(
        float(df["Q_heiz"].iloc[0] + df["Q_tww"].iloc[0]), rel=1e-9)
    # letzte Datenzeile
    assert ws.cell(row=8761, column=1).value == df.index[-1].to_pydatetime()
    # Zahlenformate (openpyxl speichert en-US-Codes; DE-Excel zeigt #.##0,0)
    assert ws.cell(row=2, column=2).number_format == "#,##0.0"
    assert "DD.MM" in ws.cell(row=2, column=1).number_format.upper()
    # Energieerhaltung: Summe der Stundenwerte = Jahresarbeit
    summe = sum(ws.cell(row=r, column=4).value for r in range(2, 8762))
    assert summe == pytest.approx(
        float((df["Q_heiz"] + df["Q_tww"]).sum()), rel=1e-9)


def test_jdl_sortiert_und_chart(wb_bytesio, ergebnisse):
    ws = wb_bytesio["JDL"]
    assert len(ws._charts) >= 1, "JDL-Blatt ohne natives Liniendiagramm"

    werte = [ws.cell(row=r, column=2).value for r in range(4, 8764)]
    werte = [v for v in werte if v is not None]
    assert len(werte) == 8760
    assert all(werte[i] >= werte[i + 1] - 1e-9 for i in range(len(werte) - 1)), \
        "JDL ist nicht absteigend sortiert"

    q_total = ergebnisse["profile"].q_total
    assert werte[0] == pytest.approx(float(q_total.max()), rel=1e-9)
    assert werte[-1] == pytest.approx(float(q_total.min()), rel=1e-9)
    # Ränge 1..8760
    assert ws.cell(row=4, column=1).value == 1
    assert ws.cell(row=8763, column=1).value == 8760


def test_tagesprofile_chart_und_struktur(wb_bytesio, ergebnisse):
    ws = wb_bytesio["Tagesprofile"]
    assert len(ws._charts) >= 1, "Tagesprofil-Blatt ohne natives Liniendiagramm"

    text = _texte(ws)
    for tagtyp in ("Werktag", "Samstag", "Sonntag"):
        assert tagtyp in text
    assert "Jahr" in text
    for monat in ("Jan", "Jul", "Dez"):
        assert monat in text

    # Jahres-Werktagsprofil muss dem analysis-Ergebnis entsprechen
    tp = analysis.tagesprofile(ergebnisse["profile"].q_total)
    # Kopfzeilen finden
    monat_row = None
    for r in range(1, 12):
        if ws.cell(row=r, column=1).value == "Stunde":
            monat_row = r
            break
    assert monat_row is not None, "Kopfzeile 'Stunde' nicht gefunden"
    start = monat_row + 2
    assert ws.cell(row=monat_row, column=2).value == "Jahr"
    assert ws.cell(row=start, column=1).value == 0
    assert ws.cell(row=start + 23, column=1).value == 23
    erwartet = float(tp[("Jahr", "Werktag")].iloc[0])
    assert ws.cell(row=start, column=2).value == pytest.approx(erwartet,
                                                               rel=1e-9)


def test_tww_blatt_enthaelt_empfehlung(wb_bytesio, ergebnisse):
    ws = wb_bytesio["TWW-Auslegung"]
    dhw = ergebnisse["dhw"]
    text = _texte(ws)

    for verfahren in ("DIN 4708", "Profilbasiert", "Faustwert"):
        assert verfahren in text, f"Verfahren '{verfahren}' fehlt"
    assert "Empfehlung" in text
    assert dhw.nl_hinweis in text
    assert "Legionell" in text
    assert dhw.details["legionellen_hinweis"] in text
    assert dhw.details["massgebend"] in text

    assert _enthaelt_zahl(ws, dhw.v_empfehlung_l), \
        "Empfohlenes Volumen taucht im TWW-Blatt nicht auf"
    for wert in (dhw.v_din4708_l, dhw.v_profil_l, dhw.v_faust_l,
                 dhw.w_z_kwh, dhw.n_bedarf):
        assert _enthaelt_zahl(ws, wert, tol=1e-6), f"Wert {wert} fehlt"

    # Empfehlung hervorgehoben (fett + Füllung)
    treffer = [c for zeile in ws.iter_rows() for c in zeile
               if c.value == "Empfohlenes Speichervolumen [l]"]
    assert treffer, "Empfehlungszeile nicht gefunden"
    zelle = treffer[0]
    assert zelle.font.bold
    wert_zelle = ws.cell(row=zelle.row, column=zelle.column + 1)
    assert wert_zelle.value == dhw.v_empfehlung_l
    assert wert_zelle.font.bold
    assert wert_zelle.fill.fgColor.rgb not in (None, "00000000")


def test_puffer_blatt_enthaelt_kriterien_und_simulation(wb_bytesio,
                                                        ergebnisse):
    ws = wb_bytesio["Puffer-Auslegung"]
    buf = ergebnisse["buffer"]
    text = _texte(ws)

    for krit in ("Abtauung", "Taktung", "EVU-Sperrzeit", "Lastgang-Simulation"):
        assert krit in text, f"Kriterium '{krit}' fehlt"
    assert buf.massgebend in text
    assert "Empfehlung" in text
    assert "Warnungen" in text

    assert _enthaelt_zahl(ws, buf.v_empfehlung_l), \
        "Empfohlenes Puffervolumen taucht im Blatt nicht auf"
    for wert in (buf.v_abtau_l, buf.v_takt_l, buf.v_sperr_l):
        assert _enthaelt_zahl(ws, wert, tol=1e-6)

    if buf.sim is not None:
        assert "Deckungsgrad" in text
        assert "Unterdeckung" in text
        assert _enthaelt_zahl(ws, buf.sim.capacity_kwh, tol=1e-6)
        assert _enthaelt_zahl(ws, buf.sim.deckungsgrad, tol=1e-9)

    for w in buf.details.get("warnungen", []):
        assert w in text
    for h in buf.details.get("hinweise", []):
        assert h in text


def test_kennzahlen_blatt(wb_bytesio, ergebnisse):
    ws = wb_bytesio["Kennzahlen"]
    kz = ergebnisse["kennzahlen"]
    text = _texte(ws)
    # Schlüsselspalte enthält die Roh-Keys
    for key in ("q_total_kwh", "p_max_kw", "vollaststunden_h", "anteil_tww",
                "p_90_prozent_kw"):
        assert key in text, f"Kennzahl '{key}' fehlt"
    assert _enthaelt_zahl(ws, kz["q_total_kwh"], tol=1e-6)
    assert _enthaelt_zahl(ws, kz["p_max_kw"], tol=1e-9)
    # verschachteltes Dict (jahressummen) wird aufgelöst
    assert "jahressummen.Q_heiz" in text


def test_eingaben_blatt(wb_bytesio, ergebnisse):
    ws = wb_bytesio["Eingaben"]
    text = _texte(ws)
    assert "Testprojekt Wärmespeicher" in text
    assert "MFH Musterstraße 1" in text
    assert "Bürotrakt" in text
    assert "Parameter TWW-Speicher" in text
    assert "Parameter Pufferspeicher" in text
    assert "Speichertemperatur [°C]" in text
    assert "Heizleistung Wärmepumpe [kW]" in text
    assert ws["A1"].font.bold
    # Spaltenbreiten gesetzt
    assert ws.column_dimensions["A"].width and ws.column_dimensions["A"].width > 20


def test_datei_und_bytesio_identisch(tmp_path, ergebnisse):
    """Der Bericht muss mit Dateipfad genauso funktionieren wie mit BytesIO."""
    pfad = tmp_path / "bericht.xlsx"
    export_excel.write_report(str(pfad), ergebnisse["profile"],
                              ergebnisse["kennzahlen"], ergebnisse["dhw"],
                              ergebnisse["buffer"], ergebnisse["projekt"])
    assert pfad.exists() and pfad.stat().st_size > 0

    wb = load_workbook(pfad)
    assert wb.sheetnames == ERWARTETE_BLAETTER
    assert wb["Lastgang"].max_row == 8761
    assert len(wb["JDL"]._charts) >= 1
    assert len(wb["Tagesprofile"]._charts) >= 1
    assert _enthaelt_zahl(wb["TWW-Auslegung"], ergebnisse["dhw"].v_empfehlung_l)
    assert _enthaelt_zahl(wb["Puffer-Auslegung"],
                          ergebnisse["buffer"].v_empfehlung_l)

    # Auch als Path-Objekt (nicht nur str)
    pfad2 = tmp_path / "bericht2.xlsx"
    export_excel.write_report(pfad2, ergebnisse["profile"],
                              ergebnisse["kennzahlen"], ergebnisse["dhw"],
                              ergebnisse["buffer"], ergebnisse["projekt"])
    assert pfad2.exists()


def test_ohne_auslegung_nicht_berechnet(ergebnisse):
    """dhw=None / buffer=None -> Blätter mit Hinweis 'nicht berechnet'."""
    puffer = io.BytesIO()
    export_excel.write_report(puffer, ergebnisse["profile"],
                              ergebnisse["kennzahlen"], None, None,
                              projekt=None)
    puffer.seek(0)
    wb = load_workbook(puffer)
    assert wb.sheetnames == ERWARTETE_BLAETTER
    assert "nicht berechnet" in _texte(wb["TWW-Auslegung"])
    assert "nicht berechnet" in _texte(wb["Puffer-Auslegung"])
    # Lastgang/JDL bleiben vollständig
    assert wb["Lastgang"].max_row == 8761
    assert len(wb["JDL"]._charts) >= 1


def test_15min_wird_auf_stunden_resampelt():
    """15-min-Profil -> 8760 Stundenzeilen, Jahresarbeit bleibt erhalten."""
    profile = _profil("15min")
    assert len(profile.df) == 4 * 8760

    puffer = io.BytesIO()
    export_excel.write_report(puffer, profile, analysis.kennzahlen(profile),
                              None, None)
    puffer.seek(0)
    ws = load_workbook(puffer)["Lastgang"]
    assert ws.max_row == 8761

    # Energieerhaltung: Summe Stundenwerte [kWh] == Jahresarbeit des 15-min-Profils
    summe_h = sum(ws.cell(row=r, column=4).value for r in range(2, 8762))
    arbeit_15 = float(profile.q_total.sum() * profile.dt_h)
    assert summe_h == pytest.approx(arbeit_15, rel=1e-9)


def test_leerer_projekt_und_messdaten_meta(ergebnisse):
    """Messdaten-Quelle wird auf dem Eingaben-Blatt ausgewiesen."""
    profile = _profil("h")
    profile.source = "gemessen"
    profile.meta = {
        "datei": "lastgang_2025.csv",
        "quelle": "Messdaten-Import",
        "einheit_quelle": "kWh",
        "aufloesung_quelle": "15min",
        "n_stunden": 8760,
        "luecken_anzahl": 2,
        "luecken_h": 3.0,
        "warnungen": ["2 Stunden interpoliert"],
    }
    puffer = io.BytesIO()
    export_excel.write_report(puffer, profile, analysis.kennzahlen(profile),
                              ergebnisse["dhw"], ergebnisse["buffer"])
    puffer.seek(0)
    text = _texte(load_workbook(puffer)["Eingaben"])
    assert "lastgang_2025.csv" in text
    assert "Messdaten-Import" in text
    assert "2 Stunden interpoliert" in text
