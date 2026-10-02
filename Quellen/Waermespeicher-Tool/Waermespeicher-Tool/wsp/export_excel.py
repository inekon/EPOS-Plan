"""Excel-Bericht des Wärmespeicher-Tools (openpyxl, native Charts).

Einstiegspunkt ist :func:`write_report`. Die Funktion schreibt sieben Blätter:

======================  ====================================================
Blatt                   Inhalt
======================  ====================================================
``Eingaben``            Projektkopf, Gebäudeliste bzw. Messdatenquelle,
                        Parameter TWW-Speicher und Pufferspeicher
``Lastgang``            stündliche Zeitreihe (Datum, Q_heiz, Q_tww, Q_total)
``JDL``                 sortierte Jahresdauerlinie + Liniendiagramm
``Tagesprofile``        mittlere Tagesprofile (Monat × Tagtyp) + Diagramm
``TWW-Auslegung``       drei Verfahren nebeneinander, Empfehlung, Hinweise
``Puffer-Auslegung``    vier Kriterien, maßgebendes Kriterium, Simulation
``Kennzahlen``          Kennzahlen-Dict tabellarisch
======================  ====================================================

Diagramme werden als **native Excel-Charts** (``openpyxl.chart.LineChart``)
erzeugt — es werden keine Bilder eingebettet.

Ziel kann ein Dateipfad **oder** ein ``io.BytesIO`` sein (Streamlit-Download).

Hinweis zu den Zahlenformaten: openpyxl speichert Formatcodes in der
en-US-Schreibweise (``#,##0.0``). Excel zeigt sie mit deutschen
Regionaleinstellungen als ``#.##0,0`` an — Tausenderpunkt und Dezimalkomma
sind also im deutschen Excel korrekt.
"""
from __future__ import annotations

import datetime as _dt
import enum as _enum
import math
from typing import Any, Optional

import pandas as pd

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import analysis
from .models import BufferResult, DHWResult, ProfileSet

__all__ = ["write_report"]

# --------------------------------------------------------------- Formatierung

#: Zahlenformate (en-US-Codes; deutsches Excel zeigt #.##0,0 an)
FMT_0 = "#,##0"
FMT_1 = "#,##0.0"
FMT_2 = "#,##0.00"
FMT_3 = "#,##0.000"
FMT_PCT = "0.0%"
FMT_DATETIME = "DD.MM.YYYY HH:MM"
FMT_DATE = "DD.MM.YYYY"

F_TITEL = Font(bold=True, size=14, color="1F3864")
F_UNTERTITEL = Font(bold=True, size=11, color="1F3864")
F_KOPF = Font(bold=True, color="FFFFFF")
F_BOLD = Font(bold=True)
F_KLEIN = Font(size=9, italic=True, color="595959")

FILL_KOPF = PatternFill("solid", start_color="1F3864")
FILL_EMPFEHLUNG = PatternFill("solid", start_color="FFE699")
FILL_WARNUNG = PatternFill("solid", start_color="FCE4D6")
FILL_HINWEIS = PatternFill("solid", start_color="E2EFDA")
FILL_BLOCK = PatternFill("solid", start_color="D9E1F2")

_DUENN = Side(style="thin", color="BFBFBF")
BORDER_BOX = Border(left=_DUENN, right=_DUENN, top=_DUENN, bottom=_DUENN)

WRAP = Alignment(wrap_text=True, vertical="top")
LINKS = Alignment(horizontal="left", vertical="top")
#: Zahlen oben ausrichten — sonst „rutschen“ Werte in hohen Zeilen nach unten
ZAHL = Alignment(horizontal="right", vertical="top")

NICHT_BERECHNET = "nicht berechnet"

SHEETS = ["Eingaben", "Lastgang", "JDL", "Tagesprofile",
          "TWW-Auslegung", "Puffer-Auslegung", "Kennzahlen"]


# ------------------------------------------------------------------ Hilfsmittel

def _clean(value: Any) -> Any:
    """Wert in einen von openpyxl schreibbaren Typ überführen."""
    if value is None:
        return None
    if isinstance(value, bool):
        return "ja" if value else "nein"
    if isinstance(value, (int, float)):
        f = float(value)
        if math.isnan(f) or math.isinf(f):
            return None
        return int(value) if isinstance(value, int) else f
    if isinstance(value, _enum.Enum):
        return _clean(value.value)
    if isinstance(value, (_dt.datetime, _dt.date, _dt.time)):
        return value
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()
    # numpy-Skalare u. Ä.
    item = getattr(value, "item", None)
    if callable(item):
        try:
            return _clean(item())
        except Exception:  # pragma: no cover - exotische Typen
            pass
    if isinstance(value, (list, tuple)):
        return "; ".join(str(_clean(v)) for v in value)
    if isinstance(value, dict):
        return "; ".join(f"{k}: {_clean(v)}" for k, v in value.items())
    return str(value)


#: Schlüssel, die als blanke Ganzzahl (ohne Tausenderpunkt) erscheinen sollen
_KEYS_PLAIN = {"jahr", "year", "try_region", "seed", "n_we", "copies",
               "start_hh", "n_schritte"}

#: Schlüssel, die als Prozentwert erscheinen sollen
_KEYS_PROZENT = {"nutzbarer_anteil", "zuschlag", "tww_share", "ziel_deckung",
                 "deckungsgrad", "sperrzeit_anteil"}


def _auto_format(key: str, value: Any) -> Optional[str]:
    """Passendes Zahlenformat anhand des Schlüsselnamens raten."""
    if not isinstance(value, (int, float)) or isinstance(value, bool):
        return None
    k = str(key).lower()
    f = float(value)
    ganzzahlig = abs(f - round(f)) < 1e-9

    if k in _KEYS_PLAIN:
        return "0"
    if k in _KEYS_PROZENT or any(
            t in k for t in ("anteil", "deckungsgrad", "abweichung")):
        return FMT_PCT

    if k.endswith("_l") or "_l_" in k or "volumen" in k or "liter" in k:
        fmt = FMT_0
    elif "kwh" in k or k.endswith("_a") or "jahressumme" in k:
        fmt = FMT_0
    elif "kw" in k:
        fmt = FMT_1
    elif k.endswith("_h") or "stunden" in k or "anzahl" in k:
        fmt = FMT_0
    elif ganzzahlig and isinstance(value, int):
        fmt = FMT_0
    else:
        fmt = FMT_1

    # kleine Nachkommawerte nicht auf 0 runden (z. B. dt_quelle_h = 0,25)
    if fmt == FMT_0 and not ganzzahlig and abs(f) < 100.0:
        fmt = FMT_2
    return fmt


def _titel(ws, row: int, text: str, breite: int = 8) -> int:
    """Titelzeile (fett) schreiben, gibt die nächste freie Zeile zurück."""
    c = ws.cell(row=row, column=1, value=text)
    c.font = F_TITEL
    if breite > 1:
        ws.merge_cells(start_row=row, start_column=1,
                       end_row=row, end_column=breite)
    return row + 2


def _abschnitt(ws, row: int, text: str, col: int = 1, breite: int = 2) -> int:
    """Abschnittsüberschrift (fett, hinterlegt)."""
    c = ws.cell(row=row, column=col, value=text)
    c.font = F_UNTERTITEL
    c.fill = FILL_BLOCK
    for i in range(1, max(breite, 1)):
        ws.cell(row=row, column=col + i).fill = FILL_BLOCK
    return row + 1


def _kv(ws, row: int, label: str, value: Any, fmt: Optional[str] = None,
        col: int = 1, einheit: str = "") -> int:
    """Label/Wert-Zeile schreiben."""
    lc = ws.cell(row=row, column=col, value=str(label))
    lc.alignment = LINKS
    v = _clean(value)
    vc = ws.cell(row=row, column=col + 1, value=v)
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        vc.number_format = fmt or FMT_1
        vc.alignment = ZAHL
    elif isinstance(v, (_dt.datetime, _dt.date)):
        vc.number_format = FMT_DATETIME if isinstance(v, _dt.datetime) else FMT_DATE
        vc.alignment = ZAHL
    else:
        vc.alignment = WRAP
    if einheit:
        ws.cell(row=row, column=col + 2, value=einheit)
    return row + 1


def _kopfzeile(ws, row: int, labels: list, col0: int = 1) -> int:
    """Tabellenkopf (fett, weiß auf blau)."""
    for i, lab in enumerate(labels):
        c = ws.cell(row=row, column=col0 + i, value=str(lab))
        c.font = F_KOPF
        c.fill = FILL_KOPF
        c.alignment = Alignment(horizontal="center", vertical="center",
                                wrap_text=True)
        c.border = BORDER_BOX
    return row + 1


def _text_block(ws, row: int, text: str, breite: int = 8,
                fill: Optional[PatternFill] = None, col: int = 1) -> int:
    """Langtext über mehrere Spalten (Zeilenumbruch aktiv)."""
    c = ws.cell(row=row, column=col, value=str(text))
    c.alignment = WRAP
    if fill is not None:
        c.fill = fill
    end_col = col + max(breite - 1, 0)
    if end_col > col:
        ws.merge_cells(start_row=row, start_column=col,
                       end_row=row, end_column=end_col)
        if fill is not None:
            for cc in range(col + 1, end_col + 1):
                ws.cell(row=row, column=cc).fill = fill
    zeilen = max(1, int(len(str(text)) / (breite * 16) + 1))
    ws.row_dimensions[row].height = max(15, 14 * min(zeilen, 6))
    return row + 1


def _breiten(ws, breiten: dict) -> None:
    """Spaltenbreiten setzen: {1: 34, 2: 16, ...} oder {'A': 34}."""
    for col, w in breiten.items():
        letter = col if isinstance(col, str) else get_column_letter(col)
        ws.column_dimensions[letter].width = w


def _dict_block(ws, row: int, daten: dict, col: int = 1,
                labels: Optional[dict] = None,
                skip: tuple = ()) -> int:
    """Dict als Label/Wert-Liste ausgeben (verschachtelte Werte als Text)."""
    if not daten:
        return _kv(ws, row, "—", NICHT_BERECHNET, col=col)
    labels = labels or {}
    for key, val in daten.items():
        if key in skip:
            continue
        label = labels.get(key, _label(key))
        row = _kv(ws, row, label, val, fmt=_auto_format(key, val), col=col)
    return row


def _label(key: str) -> str:
    """Schlüsselname -> lesbares Label."""
    if key in LABELS:
        return LABELS[key]
    txt = str(key).replace("_", " ").strip()
    return txt[:1].upper() + txt[1:]


#: Übersetzungstabelle für Detail-/Parameter-Schlüssel
LABELS = {
    # DHW-Parameter
    "t_speicher": "Speichertemperatur [°C]",
    "t_kalt": "Kaltwassertemperatur [°C]",
    "nutzbarer_anteil": "Nutzbarer Volumenanteil [-]",
    "p_lade": "Ladeleistung WP im TWW-Betrieb [kW]",
    "zirkulation_kw": "Zirkulationsverlust (dauerhaft) [kW]",
    "zuschlag": "Sicherheits-/Bereitschaftszuschlag [-]",
    # Puffer-Parameter
    "p_wp": "Heizleistung Wärmepumpe [kW]",
    "p_wp_min": "Minimale Modulationsleistung [kW]",
    "p_bivalent": "Zweiter Erzeuger (bivalent) [kW]",
    "dt_puffer": "Nutzbare Spreizung Puffer ΔT [K]",
    "t_min_lauf_min": "Mindestlaufzeit [min]",
    "abtau_l_pro_kw": "Abtauvolumen [l/kW]",
    "sperrzeiten": "EVU-Sperrzeiten (Start h, Dauer h)",
    "ziel_deckung": "Ziel-Deckungsgrad [-]",
    # Wetter
    "try_region": "TRY-Klimaregion",
    "try_region_name": "TRY-Region (Name)",
    "year": "Bezugsjahr",
    "jahr": "Bezugsjahr",
    # Gebäude
    "name": "Bezeichnung",
    "category": "Kategorie",
    "n_we": "Wohneinheiten",
    "n_persons": "Personen je WE",
    "q_heiz_a": "Q Heizung [kWh/a]",
    "q_tww_a": "Q TWW [kWh/a]",
    "copies": "Anzahl gleicher Gebäude",
    "sigma": "Zeitversatz σ [min]",
    "bdew_branch": "BDEW-Branche",
    "tww_share": "TWW-Anteil (GHD) [-]",
    # Detail-Schlüssel Auslegung
    "delta_t_nutz_K": "Nutzbare Spreizung ΔT [K]",
    "personen_gesamt": "Personen gesamt",
    "N": "Bedarfskennzahl N [-]",
    "W_z_kWh": "Wärmebedarf Zapfperiode W_z [kWh]",
    "W_z_Wh": "Wärmebedarf Zapfperiode W_z [Wh]",
    "W_1h_kWh": "Wärmebedarf 1. Stunde W_1h [kWh]",
    "W_p_kWh": "Wärmebedarf Spitzenperiode W_p [kWh]",
    "GLF": "Gleichzeitigkeitsfaktor GLF [-]",
    "v_nutz_l": "Nutzvolumen [l]",
    "methode": "Methode",
    "defizit_max_kWh": "Max. SOC-Defizit [kWh]",
    "defizit_max_zeitpunkt": "Zeitpunkt des Maximums",
    "v_roh_l": "Volumen ohne Zuschlag [l]",
    "p_lade_kW": "Ladeleistung [kW]",
    "zirkulation_kW": "Zirkulation [kW]",
    "tagesbedarf_max_kWh": "Größter Tagesbedarf [kWh/d]",
    "jahresbedarf_kWh": "Jahresbedarf [kWh/a]",
    "p_max_last_kW": "Maximale Last [kW]",
    "p_lade_ausreichend": "Ladeleistung ausreichend",
    "p_lade_mindest_kW": "Erforderliche Mindest-Ladeleistung [kW]",
    "l_pro_person_tag_60C": "Zapfmenge [l/(Person·d)] bei 60 °C",
    "v_60C_l": "Zapfmenge gesamt [l/d] bei 60 °C",
    "tagesbedarf_kWh": "Tagesbedarf [kWh/d]",
    "v_spez_l_pro_kW": "Spezifisches Volumen [l/kW]",
    "p_wp_kW": "Heizleistung WP [kW]",
    "v_l": "Volumen [l]",
    "p_wp_min_kW": "Minimale Modulationsleistung [kW]",
    "q_kWh": "Energie [kWh]",
    "dt_puffer_K": "Spreizung ΔT [K]",
    "p_gen_kW": "Erzeugerleistung P_gen [kW]",
    "ziel_deckung": "Ziel-Deckungsgrad [-]",
    "sperrzeit_anteil": "Anteil Sperrzeit am Jahr [-]",
    "deckungsgrad": "Deckungsgrad [-]",
    "unterdeckung_kWh": "Unterdeckung [kWh/a]",
    "unterdeckung_h": "Unterdeckung [h/a]",
    "fehlende_leistung_kW": "Fehlende Leistung an der Spitze [kW]",
    "gueltig": "Ziel-Deckungsgrad erreicht",
    "status": "Status",
}


# ------------------------------------------------------------ Zeitreihen-Hilfen

def _stundenwerte(profile: ProfileSet) -> pd.DataFrame:
    """Stündlicher Lastgang (kW) mit Spalten Q_heiz, Q_tww, Q_total.

    Feinere Auflösungen (z. B. 15 min) werden energieerhaltend auf 1 h
    gemittelt: Energie = Mittelwert der Leistung × 1 h.
    """
    df = profile.df.copy()
    for col in ("Q_heiz", "Q_tww"):
        if col not in df.columns:
            df[col] = 0.0
    df = df[["Q_heiz", "Q_tww"]].astype(float)

    if isinstance(df.index, pd.DatetimeIndex):
        dt_h = profile.dt_h
        if dt_h < 1.0 - 1e-9:
            df = df.resample("h").mean()
        df = df.sort_index()
    df["Q_total"] = df["Q_heiz"] + df["Q_tww"]
    return df


# ------------------------------------------------------------------- Blätter

def _blatt_eingaben(ws, profile: ProfileSet, projekt: Optional[dict]) -> None:
    _breiten(ws, {1: 38, 2: 22, 3: 16, 4: 16, 5: 16, 6: 16, 7: 16, 8: 16})
    projekt = projekt or {}
    row = _titel(ws, 1, "Wärmespeicher-Auslegung — Eingaben")

    # --- Projektkopf -------------------------------------------------------
    row = _abschnitt(ws, row, "Projekt", breite=2)
    # app.py legt den Namen unter 'name' ab; 'projekt' kann ein Unter-Dict sein
    unter = projekt.get("projekt")
    unter = unter if isinstance(unter, dict) else {}
    name = (projekt.get("projektname") or projekt.get("name")
            or unter.get("name") or "—")
    datum = (projekt.get("datum") or projekt.get("date")
             or _dt.date.today())
    row = _kv(ws, row, "Projektname", name)
    row = _kv(ws, row, "Datum", datum)
    for key in ("bearbeiter", "kunde", "ort", "bemerkung", "beschreibung"):
        if projekt.get(key):
            row = _kv(ws, row, _label(key), projekt[key])
    row = _kv(ws, row, "Erstellt am", _dt.datetime.now())
    row += 1

    # --- Datenquelle -------------------------------------------------------
    meta = dict(profile.meta or {})
    row = _abschnitt(ws, row, "Datengrundlage", breite=2)
    row = _kv(ws, row, "Quelle", profile.source)
    row = _kv(ws, row, "Zeitliche Auflösung", profile.resolution)
    row = _kv(ws, row, "Anzahl Zeitschritte", len(profile.df), fmt=FMT_0)
    if len(profile.df):
        row = _kv(ws, row, "Zeitraum von", profile.df.index[0])
        row = _kv(ws, row, "Zeitraum bis", profile.df.index[-1])
    summen = profile.annual_sums()
    for col, val in summen.items():
        row = _kv(ws, row, f"Jahressumme {col} [kWh/a]", val, fmt=FMT_0)
    row += 1

    gebaeude = projekt.get("buildings") or meta.get("gebaeude")
    if gebaeude:
        # --- Gebäudeliste --------------------------------------------------
        row = _abschnitt(ws, row, "Gebäude", breite=8)
        felder = ["name", "category", "n_we", "n_persons", "q_heiz_a",
                  "q_tww_a", "copies", "sigma", "bdew_branch", "tww_share"]
        vorhanden = [f for f in felder
                     if any(f in dict(g) for g in gebaeude)]
        row = _kopfzeile(ws, row, [_label(f) for f in vorhanden])
        for g in gebaeude:
            g = dict(g)
            for i, f in enumerate(vorhanden):
                v = _clean(g.get(f))
                c = ws.cell(row=row, column=1 + i, value=v)
                c.border = BORDER_BOX
                if isinstance(v, (int, float)) and not isinstance(v, bool):
                    c.number_format = _auto_format(f, v) or FMT_1
            row += 1
        row += 1

    wetter = projekt.get("weather")
    if wetter:
        row = _abschnitt(ws, row, "Wetter / Klimaregion", breite=2)
        row = _dict_block(ws, row, dict(wetter))
        if meta.get("try_region_name"):
            row = _kv(ws, row, "TRY-Region (Name)", meta["try_region_name"])
        row += 1

    if profile.source != "synthetisch" or meta.get("datei"):
        # --- Messdatenquelle ----------------------------------------------
        row = _abschnitt(ws, row, "Messdaten-Import", breite=2)
        for key in ("datei", "quelle", "spalten", "einheit_quelle",
                    "aufloesung_quelle", "dt_quelle_h", "n_zeilen_roh",
                    "n_stunden", "duplikate", "luecken_anzahl", "luecken_h",
                    "luecken_interpoliert", "luecken_offen",
                    "jahressumme_kwh", "p_max_kw"):
            if key in meta:
                row = _kv(ws, row, _label(key), meta[key],
                          fmt=_auto_format(key, meta[key]))
        warn = meta.get("warnungen") or []
        for w in warn:
            row = _text_block(ws, row, f"Warnung: {w}", breite=8,
                              fill=FILL_WARNUNG)
        if meta.get("tww_split"):
            row += 1
            row = _abschnitt(ws, row, "Heizung/TWW-Split", breite=2)
            row = _dict_block(ws, row, dict(meta["tww_split"]))
        row += 1

    # --- Parameter TWW -----------------------------------------------------
    row = _abschnitt(ws, row, "Parameter TWW-Speicher", breite=2)
    if projekt.get("dhw"):
        row = _dict_block(ws, row, dict(projekt["dhw"]))
    else:
        row = _kv(ws, row, "Parameter", NICHT_BERECHNET)
    row += 1

    # --- Parameter Puffer --------------------------------------------------
    row = _abschnitt(ws, row, "Parameter Pufferspeicher", breite=2)
    if projekt.get("buffer"):
        row = _dict_block(ws, row, dict(projekt["buffer"]))
    else:
        row = _kv(ws, row, "Parameter", NICHT_BERECHNET)

    ws.freeze_panes = "A3"


def _blatt_lastgang(ws, stunden: pd.DataFrame) -> None:
    _breiten(ws, {1: 20, 2: 14, 3: 14, 4: 14})
    _kopfzeile(ws, 1, ["Datum/Zeit", "Q_heiz [kW]", "Q_tww [kW]",
                       "Q_total [kW]"])

    for ts, r in zip(stunden.index, stunden.itertuples(index=False)):
        ws.append([_clean(ts), _clean(r.Q_heiz), _clean(r.Q_tww),
                   _clean(r.Q_total)])

    letzte = ws.max_row
    for r in range(2, letzte + 1):
        ws.cell(row=r, column=1).number_format = FMT_DATETIME
        for c in (2, 3, 4):
            ws.cell(row=r, column=c).number_format = FMT_1
    ws.freeze_panes = "A2"


def _blatt_jdl(ws, stunden: pd.DataFrame) -> None:
    _breiten(ws, {1: 12, 2: 16})
    row = _titel(ws, 1, "Jahresdauerlinie (Gesamtwärmeleistung)", breite=2)
    kopf = row
    row = _kopfzeile(ws, row, ["Rang [h]", "Leistung [kW]"])

    linie = analysis.jdl(stunden["Q_total"])
    start = row
    for rang, wert in zip(linie.index, linie.to_numpy()):
        ws.cell(row=row, column=1, value=int(rang)).number_format = FMT_0
        ws.cell(row=row, column=2, value=_clean(float(wert))
                ).number_format = FMT_1
        row += 1
    ende = row - 1

    if ende >= start:
        chart = LineChart()
        chart.title = "Jahresdauerlinie"
        chart.style = 12
        chart.y_axis.title = "Leistung [kW]"
        chart.x_axis.title = "Stunden [h]"
        chart.height = 11
        chart.width = 24
        daten = Reference(ws, min_col=2, min_row=kopf, max_row=ende)
        chart.add_data(daten, titles_from_data=True)
        kats = Reference(ws, min_col=1, min_row=start, max_row=ende)
        chart.set_categories(kats)
        for s in chart.series:
            s.smooth = False
        # x-Achse mit 8760 Kategorien nicht beschriften (unleserlich)
        chart.x_axis.delete = False
        chart.x_axis.tickLblSkip = max(1, (ende - start + 1) // 12)
        chart.x_axis.tickMarkSkip = chart.x_axis.tickLblSkip
        ws.add_chart(chart, "D3")

    ws.freeze_panes = "A4"


def _blatt_tagesprofile(ws, profile: ProfileSet,
                        stunden: pd.DataFrame) -> None:
    row = _titel(ws, 1, "Mittlere Tagesprofile (Gesamtwärmeleistung) [kW]",
                 breite=6)
    row = _text_block(
        ws, row,
        "Spalten: Jahr = Mittel über alle Monate, danach je Monat; "
        "Tagtypen Werktag / Samstag / Sonntag (Feiertage zählen als Sonntag).",
        breite=8)
    row += 1

    tp = analysis.tagesprofile(stunden["Q_total"])
    monat_row = row
    tagtyp_row = row + 1
    daten_start = row + 2

    ws.cell(row=monat_row, column=1, value="Stunde").font = F_KOPF
    ws.cell(row=monat_row, column=1).fill = FILL_KOPF
    ws.cell(row=tagtyp_row, column=1, value="[h]").font = F_KOPF
    ws.cell(row=tagtyp_row, column=1).fill = FILL_KOPF

    for i, (monat, tagtyp) in enumerate(tp.columns):
        col = 2 + i
        c1 = ws.cell(row=monat_row, column=col, value=str(monat))
        c1.font = F_KOPF
        c1.fill = FILL_KOPF
        c1.alignment = Alignment(horizontal="center")
        c2 = ws.cell(row=tagtyp_row, column=col, value=str(tagtyp))
        c2.font = F_KOPF
        c2.fill = FILL_KOPF
        c2.alignment = Alignment(horizontal="center")

    for j, stunde in enumerate(tp.index):
        r = daten_start + j
        ws.cell(row=r, column=1, value=int(stunde)).number_format = FMT_0
        for i, spalte in enumerate(tp.columns):
            wert = _clean(tp.iloc[j, i])
            c = ws.cell(row=r, column=2 + i, value=wert)
            c.number_format = FMT_1
    daten_ende = daten_start + len(tp.index) - 1

    _breiten(ws, {1: 10})
    for i in range(len(tp.columns)):
        _breiten(ws, {2 + i: 11})

    # Diagramm: nur die Jahres-Spalten (Werktag / Samstag / Sonntag)
    jahr_cols = [2 + i for i, (m, _t) in enumerate(tp.columns) if m == "Jahr"]
    if jahr_cols and daten_ende >= daten_start:
        chart = LineChart()
        chart.title = "Mittlere Tagesprofile (Jahr)"
        chart.style = 12
        chart.y_axis.title = "Leistung [kW]"
        chart.x_axis.title = "Stunde des Tages"
        chart.height = 10
        chart.width = 20
        daten = Reference(ws, min_col=min(jahr_cols), max_col=max(jahr_cols),
                          min_row=tagtyp_row, max_row=daten_ende)
        chart.add_data(daten, titles_from_data=True)
        kats = Reference(ws, min_col=1, min_row=daten_start,
                         max_row=daten_ende)
        chart.set_categories(kats)
        anker_col = get_column_letter(len(tp.columns) + 3)
        ws.add_chart(chart, f"{anker_col}{daten_start}")

    ws.freeze_panes = ws.cell(row=daten_start, column=2).coordinate


def _blatt_tww(ws, dhw: Optional[DHWResult]) -> None:
    _breiten(ws, {1: 40, 2: 18, 3: 3, 4: 40, 5: 18, 6: 3, 7: 40, 8: 18})
    row = _titel(ws, 1, "Trinkwarmwasser-Speicher — Auslegung")

    if dhw is None:
        _text_block(ws, row,
                    f"TWW-Speicher: {NICHT_BERECHNET}. "
                    "Es wurden keine Auslegungsparameter übergeben bzw. die "
                    "Berechnung wurde nicht durchgeführt.",
                    breite=8, fill=FILL_WARNUNG)
        return

    det = dhw.details or {}

    # --- Vergleich der Verfahren ------------------------------------------
    row = _abschnitt(ws, row, "Vergleich der Verfahren", breite=8)
    row = _kopfzeile(ws, row, ["Verfahren", "Volumen [l]", "", "Kurzbeschreibung"])
    massgebend = det.get("massgebend", "—")
    verfahren = [
        ("DIN 4708", dhw.v_din4708_l,
         "Bedarfskennzahl N, Wärmebedarf der Zapfperiode W_z (Norm-Reserven "
         "bereits enthalten)"),
        ("Profilbasiert", dhw.v_profil_l,
         "SOC-Defizit aus dem TWW-Lastgang bei konstanter Ladeleistung "
         "(inkl. Zuschlag)"),
        ("Faustwert", dhw.v_faust_l,
         "35 l/(Person·d) bei 60 °C, energiegleich umgerechnet (inkl. Zuschlag)"),
    ]
    for name, vol, beschr in verfahren:
        c0 = ws.cell(row=row, column=1, value=name)
        c0.alignment = LINKS
        c1 = ws.cell(row=row, column=2, value=_clean(vol))
        c1.number_format = FMT_1
        c1.alignment = ZAHL
        ws.cell(row=row, column=4, value=beschr).alignment = WRAP
        for c in (c0, c1):
            c.border = BORDER_BOX
        if name == massgebend:
            c0.font = F_BOLD
            c1.font = F_BOLD
            c0.fill = FILL_HINWEIS
            c1.fill = FILL_HINWEIS
        if not vol:
            ws.cell(row=row, column=3, value="entfällt").font = F_KLEIN
        row += 1
    row += 1

    # --- Empfehlung (hervorgehoben) ---------------------------------------
    row = _abschnitt(ws, row, "Empfehlung", breite=8)
    for label, wert, fmt, breit in (
            ("Maßgebendes Verfahren", massgebend, None, False),
            ("Erforderliches Volumen (ungerundet) [l]",
             det.get("v_max_ungerundet_l"), FMT_1, False),
            ("Empfohlenes Speichervolumen [l]", dhw.v_empfehlung_l, FMT_0,
             False),
            ("Hinweis Leistungskennzahl", dhw.nl_hinweis, None, True)):
        lc = ws.cell(row=row, column=1, value=label)
        lc.alignment = LINKS
        vc = ws.cell(row=row, column=2, value=_clean(wert))
        if fmt:
            vc.number_format = fmt
            vc.alignment = ZAHL
        else:
            vc.alignment = LINKS
        for c in (lc, vc):
            c.font = F_BOLD
            c.fill = FILL_EMPFEHLUNG
            c.border = BORDER_BOX
        if breit:  # Langtext über die Spalten B..H führen
            ws.merge_cells(start_row=row, start_column=2,
                           end_row=row, end_column=8)
            for cc in range(3, 9):
                ws.cell(row=row, column=cc).fill = FILL_EMPFEHLUNG
        row += 1
    row += 1

    # --- Verfahren nebeneinander (Zwischenwerte) --------------------------
    row = _abschnitt(ws, row, "Zwischenwerte der Verfahren", breite=8)
    block_row = row
    bloecke = [
        (1, "1) DIN 4708", det.get("din4708") or {},
         [("Speichervolumen [l]", dhw.v_din4708_l, FMT_1)]),
        (4, "2) Profilbasiert", det.get("profil") or {},
         [("Speichervolumen [l]", dhw.v_profil_l, FMT_1)]),
        (7, "3) Faustwert", det.get("faustwert") or {},
         [("Speichervolumen [l]", dhw.v_faust_l, FMT_1)]),
    ]
    ende_zeilen = []
    for col, titel, daten, kopfwerte in bloecke:
        r = block_row
        c = ws.cell(row=r, column=col, value=titel)
        c.font = F_UNTERTITEL
        c.fill = FILL_BLOCK
        ws.cell(row=r, column=col + 1).fill = FILL_BLOCK
        r += 1
        for label, wert, fmt in kopfwerte:
            r = _kv(ws, r, label, wert, fmt=fmt, col=col)
        if daten:
            r = _dict_block(ws, r, dict(daten), col=col, skip=("methode",))
            if daten.get("methode"):
                mc = ws.cell(row=r, column=col, value=str(daten["methode"]))
                mc.font = F_KLEIN
                mc.alignment = WRAP
                ws.merge_cells(start_row=r, start_column=col,
                               end_row=r, end_column=col + 1)
                ws.row_dimensions[r].height = 46
                r += 1
        else:
            r = _kv(ws, r, "Status", "Verfahren entfällt (kein Lastgang)",
                    col=col)
        ende_zeilen.append(r)
    row = max(ende_zeilen) + 1

    # --- Rahmenparameter ---------------------------------------------------
    row = _abschnitt(ws, row, "Rahmenbedingungen", breite=8)
    for key in ("delta_t_nutz_K", "nutzbarer_anteil", "zuschlag",
                "personen_gesamt"):
        if key in det:
            row = _kv(ws, row, _label(key), det[key],
                      fmt=_auto_format(key, det[key]))
    row += 1

    # --- Legionellen / Warnungen / Hinweise --------------------------------
    row = _abschnitt(ws, row, "Legionellenschutz (DVGW W 551 / VDI 6023)",
                     breite=8)
    row = _text_block(ws, row, det.get("legionellen_hinweis", "—"),
                      breite=8, fill=FILL_HINWEIS)
    row += 1

    warnungen = det.get("warnungen") or []
    row = _abschnitt(ws, row, "Warnungen", breite=8)
    if warnungen:
        for w in warnungen:
            row = _text_block(ws, row, f"⚠ {w}", breite=8, fill=FILL_WARNUNG)
    else:
        row = _text_block(ws, row, "Keine Warnungen.", breite=8)
    row += 1

    hinweise = det.get("hinweise") or []
    row = _abschnitt(ws, row, "Hinweise", breite=8)
    if hinweise:
        for h in hinweise:
            row = _text_block(ws, row, f"• {h}", breite=8)
    else:
        row = _text_block(ws, row, "Keine Hinweise.", breite=8)

    ws.freeze_panes = "A3"


def _sim_methode(details: dict) -> str:
    """Kurzbeschreibung des Simulationskriteriums inkl. Status."""
    sim = details.get("simulation") or {}
    txt = ("Binärsuche über die Kapazität: SOC-Bilanz über den Lastgang mit "
           "P_gen = P_WP + P_bivalent und Sperrzeiten als Verfügbarkeitsmaske")
    status = str(sim.get("status", "")).strip()
    if status and status.lower() != "ok":
        txt += f" — {status}"
    if sim.get("gueltig") is False:
        txt += (" — Ziel-Deckungsgrad nicht erreicht, Kriterium geht NICHT in "
                "die Empfehlung ein")
    return txt


def _blatt_puffer_betrieb(ws, det: dict, row: int) -> int:
    """Block „Betriebs-Simulation (Zweipunkt)" — nur wenn Kennwerte vorliegen.

    Erwartet ``det['betriebssimulation']`` als dict mit den Schlüsseln
    ``volumen_l``, ``kapazitaet_kwh``, ``min_soc_frac``, ``ladezyklen``,
    ``unterdeckung_kwh``, ``unterdeckung_h``, ``deckungsgrad``,
    ``laufzeit_mittel_h``. Fehlt der Eintrag oder ist er unvollständig, wird
    der Block übersprungen bzw. mit den vorhandenen Werten geschrieben.
    """
    betrieb = (det or {}).get("betriebssimulation")
    if not isinstance(betrieb, dict) or not betrieb:
        return row

    row = _abschnitt(ws, row, "Betriebs-Simulation (Zweipunkt, SWSG-Logik)",
                     breite=8)
    row = _text_block(
        ws, row,
        "Zweipunkt-Betrieb eines konkret gewählten Speichers: Der Erzeuger "
        "ist AN (deckt die Last und lädt den Speicher bis voll) oder AUS "
        "(die Last wird allein aus dem Speicher gedeckt, bis der "
        "Mindestfüllstand erreicht ist). Im Gegensatz zur Auslegungs-"
        "Simulation (Frage: erforderliche Größe) zeigt sie das "
        "Betriebsverhalten — insbesondere die Taktung.", breite=8)

    felder = (
        ("Gewähltes Speichervolumen [l]", "volumen_l", FMT_1),
        ("Nutzbare Kapazität [kWh]", "kapazitaet_kwh", FMT_2),
        ("Minimaler Füllstand [-]", "min_soc_frac", FMT_PCT),
        ("Ladezyklen [1/a]", "ladezyklen", FMT_0),
        ("Mittlere Laufzeit je Zyklus [h]", "laufzeit_mittel_h", FMT_2),
        ("Unterdeckung [kWh/a]", "unterdeckung_kwh", FMT_1),
        ("Stunden mit Unterdeckung [h/a]", "unterdeckung_h", FMT_0),
        ("Deckungsgrad [-]", "deckungsgrad", FMT_PCT),
    )
    for label, key, fmt in felder:
        if key in betrieb:
            row = _kv(ws, row, label, betrieb.get(key), fmt)
    # unbekannte Zusatzschlüssel nicht verschlucken
    rest = {k: v for k, v in betrieb.items()
            if k not in {key for _, key, _ in felder}}
    if rest:
        row = _dict_block(ws, row, rest)
    return row + 1


def _blatt_puffer(ws, buffer: Optional[BufferResult],
                  dt_h: float = 1.0) -> None:
    _breiten(ws, {1: 44, 2: 18, 3: 3, 4: 46, 5: 18, 6: 16, 7: 16, 8: 16})
    row = _titel(ws, 1, "Pufferspeicher — Auslegung")

    if buffer is None:
        _text_block(ws, row,
                    f"Pufferspeicher: {NICHT_BERECHNET}. "
                    "Es wurden keine Auslegungsparameter übergeben bzw. die "
                    "Berechnung wurde nicht durchgeführt.",
                    breite=8, fill=FILL_WARNUNG)
        return

    det = buffer.details or {}

    # --- Vier Kriterien ----------------------------------------------------
    row = _abschnitt(ws, row, "Kriterien im Vergleich", breite=8)
    row = _kopfzeile(ws, row, ["Kriterium", "Volumen [l]", "", "Methode"])
    kriterien = [
        ("Abtauung", buffer.v_abtau_l,
         (det.get("abtauung") or {}).get("methode", "")),
        ("Taktung", buffer.v_takt_l,
         (det.get("taktung") or {}).get("methode", "")),
        ("EVU-Sperrzeit", buffer.v_sperr_l,
         (det.get("sperrzeit") or {}).get(
             "hinweis", "V aus maximaler Mittelleistung im Sperrfenster")),
        ("Lastgang-Simulation", buffer.v_sim_l, _sim_methode(det)),
    ]
    for name, vol, methode in kriterien:
        c0 = ws.cell(row=row, column=1, value=name)
        c0.alignment = LINKS
        c1 = ws.cell(row=row, column=2, value=_clean(vol))
        c1.number_format = FMT_1
        c1.alignment = ZAHL
        ws.cell(row=row, column=4, value=str(methode)).alignment = WRAP
        for c in (c0, c1):
            c.border = BORDER_BOX
        if name == buffer.massgebend:
            c0.font = F_BOLD
            c1.font = F_BOLD
            c0.fill = FILL_HINWEIS
            c1.fill = FILL_HINWEIS
        if not vol:
            ws.cell(row=row, column=3, value="entfällt").font = F_KLEIN
        row += 1
    row += 1

    # --- Empfehlung --------------------------------------------------------
    row = _abschnitt(ws, row, "Empfehlung", breite=8)
    for label, wert, fmt in (
            ("Maßgebendes Kriterium", buffer.massgebend, None),
            ("Erforderliches Volumen (ungerundet) [l]",
             det.get("v_max_ungerundet_l"), FMT_1),
            ("Empfohlenes Puffervolumen [l]", buffer.v_empfehlung_l, FMT_0)):
        lc = ws.cell(row=row, column=1, value=label)
        vc = ws.cell(row=row, column=2, value=_clean(wert))
        if fmt:
            vc.number_format = fmt
        for c in (lc, vc):
            c.font = F_BOLD
            c.fill = FILL_EMPFEHLUNG
            c.border = BORDER_BOX
        row += 1
    row += 1

    # --- Zwischenwerte je Kriterium ---------------------------------------
    for key, titel in (("abtauung", "1) Abtauung"),
                       ("taktung", "2) Taktung"),
                       ("sperrzeit", "3) EVU-Sperrzeit"),
                       ("simulation", "4) Lastgang-Simulation")):
        daten = det.get(key)
        row = _abschnitt(ws, row, titel, breite=2)
        if not daten:
            row = _kv(ws, row, "Status", NICHT_BERECHNET)
            row += 1
            continue
        daten = dict(daten)
        fenster = daten.pop("fenster", None)
        methode = daten.pop("methode", None)
        row = _dict_block(ws, row, daten)
        if fenster:
            row = _kopfzeile(ws, row, ["Sperrfenster (Start h)", "Dauer [h]",
                                       "", "Mittelleistung [kW]",
                                       "Energie [kWh]", "Volumen [l]"])
            for f in fenster:
                ws.cell(row=row, column=1, value=_clean(f.get("start_hh"))
                        ).number_format = FMT_1
                ws.cell(row=row, column=2, value=_clean(f.get("dauer_h"))
                        ).number_format = FMT_1
                ws.cell(row=row, column=4, value=_clean(f.get("q_mittel_kW"))
                        ).number_format = FMT_1
                ws.cell(row=row, column=5, value=_clean(f.get("q_kWh"))
                        ).number_format = FMT_1
                ws.cell(row=row, column=6, value=_clean(f.get("v_l"))
                        ).number_format = FMT_1
                ws.cell(row=row, column=7, value=str(f.get("quelle", "")))
                row += 1
        if methode:
            mc = ws.cell(row=row, column=1, value=str(methode))
            mc.font = F_KLEIN
            mc.alignment = WRAP
            row += 1
        row += 1

    # --- Simulationskennwerte ---------------------------------------------
    row = _abschnitt(ws, row, "Kennwerte der Speichersimulation", breite=8)
    sim = buffer.sim
    if sim is None:
        row = _text_block(ws, row,
                          "Keine Lastgang-Simulation vorhanden "
                          f"({NICHT_BERECHNET}).", breite=8)
    else:
        row = _kv(ws, row, "Nutzbare Kapazität [kWh]", sim.capacity_kwh, FMT_2)
        row = _kv(ws, row, "Zugehöriges Volumen [l]", buffer.v_sim_l, FMT_1)
        row = _kv(ws, row, "Kapazität lt. Ergebnis c_sim [kWh]",
                  buffer.c_sim_kwh, FMT_2)
        row = _kv(ws, row, "Unterdeckung [kWh/a]", sim.unterdeckung_kwh, FMT_1)
        row = _kv(ws, row, "Stunden mit Unterdeckung [h/a]",
                  sim.unterdeckung_h, FMT_0)
        row = _kv(ws, row, "Deckungsgrad [-]", sim.deckungsgrad, FMT_PCT)
        try:
            q_last = sim.df["Q_last"].astype(float)
            row = _kv(ws, row, "Jahreswärmebedarf (simuliert) [kWh/a]",
                      float(q_last.sum() * dt_h), FMT_0)
            row = _kv(ws, row, "Maximale Last [kW]", float(q_last.max()), FMT_1)
            row = _kv(ws, row, "Maximaler Speicherfüllstand [kWh]",
                      float(sim.df["SOC"].max()), FMT_2)
            row = _kv(ws, row, "Minimaler Speicherfüllstand [kWh]",
                      float(sim.df["SOC"].min()), FMT_2)
        except Exception:  # pragma: no cover - defensive
            pass
    row += 1

    # --- Betriebs-Simulation (Zweipunkt) -----------------------------------
    row = _blatt_puffer_betrieb(ws, det, row)

    # --- Warnungen / Hinweise ---------------------------------------------
    warnungen = det.get("warnungen") or []
    row = _abschnitt(ws, row, "Warnungen", breite=8)
    if warnungen:
        for w in warnungen:
            row = _text_block(ws, row, f"⚠ {w}", breite=8, fill=FILL_WARNUNG)
    else:
        row = _text_block(ws, row, "Keine Warnungen.", breite=8)
    row += 1

    hinweise = det.get("hinweise") or []
    row = _abschnitt(ws, row, "Hinweise", breite=8)
    if hinweise:
        for h in hinweise:
            row = _text_block(ws, row, f"• {h}", breite=8)
    else:
        row = _text_block(ws, row, "Keine Hinweise.", breite=8)

    ws.freeze_panes = "A3"


def _blatt_kennzahlen(ws, kennzahlen: Optional[dict]) -> None:
    _breiten(ws, {1: 46, 2: 20, 3: 26})
    row = _titel(ws, 1, "Kennzahlen des Lastgangs", breite=3)

    if not kennzahlen:
        _text_block(ws, row, f"Kennzahlen: {NICHT_BERECHNET}.", breite=3,
                    fill=FILL_WARNUNG)
        return

    row = _kopfzeile(ws, row, ["Kennzahl", "Wert", "Schlüssel"])
    start = row
    for key, val in kennzahlen.items():
        if isinstance(val, dict):
            for k2, v2 in val.items():
                lc = ws.cell(row=row, column=1,
                             value=f"{_label(key)} — {_label(k2)}")
                vc = ws.cell(row=row, column=2, value=_clean(v2))
                if isinstance(_clean(v2), (int, float)):
                    vc.number_format = _auto_format(f"{key}_{k2}", v2) or FMT_1
                ws.cell(row=row, column=3, value=f"{key}.{k2}").font = F_KLEIN
                lc.alignment = LINKS
                row += 1
            continue
        lc = ws.cell(row=row, column=1, value=_label(key))
        lc.alignment = LINKS
        v = _clean(val)
        vc = ws.cell(row=row, column=2, value=v)
        if isinstance(v, (int, float)) and not isinstance(v, bool):
            vc.number_format = _auto_format(key, val) or FMT_1
        elif isinstance(v, _dt.datetime):
            vc.number_format = FMT_DATETIME
        else:
            vc.alignment = WRAP
        ws.cell(row=row, column=3, value=str(key)).font = F_KLEIN
        row += 1

    for r in range(start, row):
        for c in (1, 2, 3):
            ws.cell(row=r, column=c).border = BORDER_BOX

    ws.freeze_panes = f"A{start}"


# ------------------------------------------------------------------ Hauptroutine

def write_report(path_or_buffer,
                 profile: ProfileSet,
                 kennzahlen: dict,
                 dhw: Optional[DHWResult],
                 buffer: Optional[BufferResult],
                 projekt: Optional[dict] = None) -> None:
    """Schreibt den vollständigen Excel-Bericht.

    Args:
        path_or_buffer: Dateipfad (str/Path) **oder** dateiähnliches Objekt
            (z. B. ``io.BytesIO`` für den Streamlit-Download).
        profile: Lastgang (:class:`wsp.models.ProfileSet`). Feinere
            Auflösungen als 1 h werden für das Blatt ``Lastgang``
            energieerhaltend auf Stundenwerte gemittelt.
        kennzahlen: Ergebnis von :func:`wsp.analysis.kennzahlen`.
        dhw: Ergebnis von :func:`wsp.sizing_dhw.size_dhw` oder ``None``
            (Blatt wird mit dem Hinweis „nicht berechnet“ angelegt).
        buffer: Ergebnis von :func:`wsp.sizing_buffer.size_buffer` oder
            ``None``.
        projekt: Projekt-Dict; enthält üblicherweise die per
            :func:`wsp.models.project_to_dict` serialisierten Eingaben
            (``buildings``, ``weather``, ``dhw``, ``buffer``) sowie
            optional ``name``/``projektname``, ``datum``, ``bearbeiter``.

    Returns:
        None — die Datei wird geschrieben bzw. der Puffer befüllt.
    """
    if profile is None:
        raise ValueError("profile darf nicht None sein")

    stunden = _stundenwerte(profile)

    wb = Workbook()
    ws_eingaben = wb.active
    ws_eingaben.title = "Eingaben"
    ws_lastgang = wb.create_sheet("Lastgang")
    ws_jdl = wb.create_sheet("JDL")
    ws_tp = wb.create_sheet("Tagesprofile")
    ws_tww = wb.create_sheet("TWW-Auslegung")
    ws_puffer = wb.create_sheet("Puffer-Auslegung")
    ws_kennzahlen = wb.create_sheet("Kennzahlen")

    _blatt_eingaben(ws_eingaben, profile, projekt)
    _blatt_lastgang(ws_lastgang, stunden)
    _blatt_jdl(ws_jdl, stunden)
    _blatt_tagesprofile(ws_tp, profile, stunden)
    _blatt_tww(ws_tww, dhw)
    _blatt_puffer(ws_puffer, buffer, dt_h=profile.dt_h)
    _blatt_kennzahlen(ws_kennzahlen, kennzahlen)

    wb.save(path_or_buffer)
